"""Источник excel_file: факт маркетинга из Excel-ведомости (Ф4, срез 1).

Схема «как у клиента»: план приходит из одного места (пока — тестовый датасет,
позже Google-таблица), факт — из Excel-ведомости, присланной файлом. Источник
читает .xlsx по EXCEL_PATH (или data/fact_marketing.xlsx) через openpyxl,
сверяет каналы и месяцы с планом и подменяет факт в данных тестового датасета.

Остальные секции (ДДС, продажи) наследуются от MockSource без изменений —
их источники (FinTablo/Битрикс) подключаются следующими срезами Ф4.
"""
import os
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook

from .mock_source import MockSource
from .sources import SourceError

APP_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.dirname(os.path.dirname(APP_DIR))
DEFAULT_XLSX = os.path.join("data", "fact_marketing.xlsx")
SHEET_TITLE = u"Факт маркетинга"
HEADER_ROW = 2


def _money(value: Any) -> int:
    try:
        return int(round(float(str(value).replace(" ", "").replace(u"\xa0", "").replace(",", "."))))
    except (TypeError, ValueError):
        raise SourceError(u"в ведомости не число: %r" % (value,))


class ExcelSource(MockSource):
    """Факт маркетинга — из Excel; план и остальные секции — из тестового датасета."""

    name = "excel_file"
    title = u"Excel-файл (факт маркетинга)"

    def __init__(self, data_file=None, xlsx_file=None) -> None:
        MockSource.__init__(self, data_file)
        raw = xlsx_file or os.environ.get("EXCEL_PATH") or DEFAULT_XLSX
        self.xlsx_file = raw if os.path.isabs(raw) else os.path.join(BASE_DIR, raw)
        self._xlsx_cache = None  # type: Optional[Dict[str, List[int]]]
        self._xlsx_header_cache = None  # type: Optional[List[str]]
        self._xlsx_key = None  # type: Optional[float]

    # ------------------------------------------------------------- ведомость --
    def _workbook_rows(self):
        try:
            key = os.path.getmtime(self.xlsx_file)
        except OSError:
            raise SourceError(u"Excel-ведомость не найдена: %s" % self.xlsx_file)
        if self._xlsx_cache is None or self._xlsx_key != key:
            try:
                wb = load_workbook(self.xlsx_file, read_only=True, data_only=True)
            except Exception as exc:
                raise SourceError(u"не удалось открыть ведомость: %s" % exc)
            ws = wb[SHEET_TITLE] if SHEET_TITLE in wb.sheetnames else wb.active
            rows = list(ws.iter_rows(values_only=True))
            wb.close()
            if len(rows) < HEADER_ROW + 1:
                raise SourceError(u"ведомость пуста: %s" % self.xlsx_file)
            header = [str(c).strip() if c is not None else "" for c in rows[HEADER_ROW - 1]]
            months = header[1:]
            if not months:
                raise SourceError(u"в шапке ведомости нет месяцев")
            fact = {}
            for row in rows[HEADER_ROW:]:
                if row is None or not row or row[0] is None:
                    continue
                channel = str(row[0]).strip()
                values = list(row[1:len(months) + 1])
                if len(values) < len(months):
                    values += [None] * (len(months) - len(values))
                fact[channel] = [_money(v) for v in values]
            if not fact:
                raise SourceError(u"в ведомости нет строк с каналами")
            self._xlsx_cache = fact
            self._xlsx_header_cache = months
            self._xlsx_key = key
        return self._xlsx_cache, self._xlsx_header_cache

    def rows_count(self) -> int:
        fact, _ = self._workbook_rows()
        return MockSource.rows_count(self) + len(fact)

    # ------------------------------------------------------------------ API --
    def _merged_dataset(self):
        """Датасет с фактом, подменённым из ведомости. Точка расширения:
        наследники (google_sheets) дополнительно подменяют план до отдачи."""
        fact, months_xlsx = self._workbook_rows()
        d = self._dataset()
        channels = d.get("marketing", {}).get("channels", [])

        plan_names = [ch.get("channel", "") for ch in channels]
        unknown = sorted(set(fact) - set(plan_names))
        missing = [n for n in plan_names if n not in fact]
        problems = []
        if unknown:
            problems.append(u"нет в плане: %s" % u", ".join(unknown))
        if missing:
            problems.append(u"нет в ведомости: %s" % u", ".join(missing))
        if problems:
            raise SourceError(u"ведомость не сходится с планом (%s)" % u"; ".join(problems))
        if months_xlsx and list(months_xlsx) != list(d.get("months", [])):
            raise SourceError(u"месяцы ведомости (%s) не совпадают с датасетом (%s)"
                              % (u", ".join(months_xlsx), u", ".join(d.get("months", []))))

        merged = dict(d)
        merged_channels = []
        for ch in channels:
            ch_copy = dict(ch)
            ch_copy["fact"] = list(fact[ch.get("channel", "")])
            merged_channels.append(ch_copy)
        merged["marketing"] = {"channels": merged_channels}
        return merged

    def fetch(self, period):
        """Факт из ведомости подменяет факт датасета — дальше всё считает MockSource."""
        if period not in ("all", "q3", "sep"):
            raise SourceError(u"неизвестный период: %s" % period)
        return self._render(self._merged_dataset(), period)

    def _render(self, merged, period):
        """Считает fetch() от объединённого датасета: _dataset() временно
        возвращает merged, иначе он перечитает JSON с диска и затрёт подмену."""
        try:
            self._dataset = lambda: merged  # instance-атрибут перекрывает метод
            return MockSource.fetch(self, period)
        finally:
            del self._dataset  # вернуть метод класса

    def fact_file(self) -> str:
        """Имя файла ведомости — для журнала импортов."""
        return os.path.basename(self.xlsx_file)
