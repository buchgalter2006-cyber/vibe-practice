"""Точка входа каркаса: Flask отдаёт статику из app/static и JSON API.

Запуск:  .venv/bin/python app/server.py   ->  http://127.0.0.1:5000
Порт берётся из .env (PORT, по умолчанию 5000), источник данных — DATA_SOURCE.
Python 3.9: без match, без X | Y в аннотациях.
"""
import functools
import io
import math
import os
import sys
import tempfile
import traceback
from typing import Any, Dict, List

APP_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.dirname(APP_DIR)
for _p in (BASE_DIR, APP_DIR):  # чтобы работало и `python app/server.py`, и `python -m app.server`
    if _p not in sys.path:
        sys.path.insert(0, _p)

from dotenv import load_dotenv  # noqa: E402

load_dotenv(os.path.join(BASE_DIR, ".env"))

from flask import Flask, jsonify, redirect, request, send_from_directory, session  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from werkzeug.security import check_password_hash, generate_password_hash  # noqa: E402

import db  # noqa: E402
from services.excel_source import SHEET_TITLE, ExcelSource  # noqa: E402
from services.sources import (  # noqa: E402
    PERIODS,
    PERIOD_LABELS,
    SourceError,
    available_sources,
    get_source,
)

APP_NAME = u"Сервис отчётности «Вектор» — каркас (Ф3)"


class ApiError(Exception):
    """Ошибка API: наружу уходит {"error": "..."} + HTTP-код."""

    def __init__(self, message: str, code: int = 400) -> None:
        super(ApiError, self).__init__(message)
        self.message = message
        self.code = code


app = Flask(__name__, static_folder="static", static_url_path="")
app.json.ensure_ascii = False   # кириллица в ответах — как есть, UTF-8
app.json.sort_keys = False
app.config["APP_NAME"] = APP_NAME
# Ф5: сессии (cookie) — секрет из .env; без него сессии не переживают перезапуск
app.secret_key = os.environ.get("SESSION_SECRET") or os.urandom(32)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024  # загружаемые ведомости маленькие

SOURCE = None  # type: Any  # инициализируется в init_app()

# Ф5: роли доступа (чем выше число — тем шире права)
ROLE_LEVELS = {"view": 1, "editor": 2, "admin": 3}
ROLES = ("admin", "editor", "view")
# Порог контрольной сверки загружаемой ведомости с текущей (доля расхождения)
FACT_DEVIATION_LIMIT = 0.5
# Хэширование паролей: pbkdf2 явно — scrypt в этом Python 3.9 недоступен
PASSWORD_METHOD = "pbkdf2:sha256:600000"


# --------------------------------------------------------------------------- #
# инфраструктура
# --------------------------------------------------------------------------- #
def get_active_source() -> Any:
    global SOURCE
    if SOURCE is None:
        SOURCE = get_source()
    return SOURCE


def data(period: str) -> Dict[str, Any]:
    """Единая точка получения данных от активного источника."""
    return get_active_source().fetch(period)


def parse_period() -> str:
    period = (request.args.get("period") or "all").strip().lower()
    if period not in PERIODS:
        raise ApiError(u"неизвестный период: «%s» (доступно: %s)" % (period, ", ".join(PERIODS)))
    return period


@app.errorhandler(ApiError)
def handle_api_error(exc):
    return jsonify({"error": exc.message}), exc.code


@app.errorhandler(SourceError)
def handle_source_error(exc):
    return jsonify({"error": str(exc)}), 503


@app.errorhandler(404)
def handle_404(exc):
    if request.path.startswith("/api/"):
        return jsonify({"error": u"нет такого метода API: %s" % request.path}), 404
    return jsonify({"error": u"не найдено: %s" % request.path}), 404


@app.errorhandler(Exception)
def handle_unexpected(exc):
    # ApiError/SourceError перехватываются своими обработчиками выше
    traceback.print_exc()
    if request.path.startswith("/api/"):
        return jsonify({"error": u"внутренняя ошибка сервера: %s" % exc.__class__.__name__}), 500
    return jsonify({"error": u"внутренняя ошибка сервера"}), 500


# --------------------------------------------------------------------------- #
# авторизация и роли (Ф5)
# --------------------------------------------------------------------------- #
def current_user() -> Any:
    """Пользователь из сессии (или None). Существование проверяем по БД."""
    name = session.get("user")
    if not name:
        return None
    row = db.get_user(name)
    return row if row else None


def _require_user() -> Dict[str, Any]:
    """401, если в сессии нет валидного пользователя."""
    user = current_user()
    if not user:
        raise ApiError(u"требуется вход", 401)
    return user


def require_role(min_role: str):
    """Декоратор: не залогинен → 401, роль ниже нужной → 403."""
    def decorator(fn):
        @functools.wraps(fn)
        def wrapper(*args, **kwargs):
            user = _require_user()
            if ROLE_LEVELS.get(user["role"], 0) < ROLE_LEVELS[min_role]:
                raise ApiError(u"недостаточно прав (нужно: %s)" % min_role, 403)
            return fn(*args, **kwargs)
        return wrapper
    return decorator


@app.route("/login")
def login_page():
    return send_from_directory(app.static_folder, "login.html")


@app.route("/login", methods=["POST"])
def api_login():
    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        payload = {"username": request.form.get("username"),
                   "password": request.form.get("password")}
    username = (payload.get("username") or "").strip()
    password = payload.get("password") or ""
    user = db.get_user(username)
    if not user or not check_password_hash(user["password_hash"], password):
        if not request.is_json:
            return redirect("/login?error=1")
        raise ApiError(u"неверный логин или пароль", 401)
    session.clear()
    session["user"] = username
    if not request.is_json:
        return redirect("/#overview")
    return jsonify({"ok": True, "user": {"username": username, "role": user["role"]}})


@app.route("/logout", methods=["POST"])
def api_logout():
    session.clear()
    return jsonify({"ok": True})


@app.route("/api/whoami")
def api_whoami():
    user = current_user()
    if not user:
        return jsonify({"user": None, "role": None})
    return jsonify({"user": user["username"], "role": user["role"]})


# --------------------------------------------------------------------------- #
# управление пользователями (только admin; смена своего пароля — любым)
# --------------------------------------------------------------------------- #
def _valid_password(password: str) -> str:
    if not password or len(password) < 8:
        raise ApiError(u"пароль должен быть не короче 8 символов")
    return password


@app.route("/api/users", methods=["GET", "POST"])
@require_role("admin")
def api_users():
    if request.method == "GET":
        return jsonify({"users": db.list_users()})

    payload = request.get_json(silent=True) or {}
    username = (payload.get("username") or "").strip()
    role = (payload.get("role") or "").strip()
    password = payload.get("password") or ""
    if not username or len(username) > 40 or any(ch.isspace() for ch in username):
        raise ApiError(u"имя пользователя: 1–40 символов без пробелов")
    if role not in ROLES:
        raise ApiError(u"роль должна быть одной из: %s" % ", ".join(ROLES))
    _valid_password(password)
    if db.get_user(username):
        raise ApiError(u"пользователь «%s» уже существует" % username)
    db.insert_user(username, generate_password_hash(password, method=PASSWORD_METHOD), role)
    return jsonify({"ok": True, "user": {"username": username, "role": role}})


@app.route("/api/users/<username>/password", methods=["POST"])
def api_user_password(username):
    """Сменить пароль: себе — любой ролью, другому — только admin."""
    me = _require_user()
    if username != me["username"] and me["role"] != "admin":
        raise ApiError(u"менять пароль другим пользователям может только admin", 403)
    payload = request.get_json(silent=True) or {}
    password = _valid_password(payload.get("password") or "")
    if not db.get_user(username):
        raise ApiError(u"пользователь не найден: %s" % username, 404)
    db.set_password(username, generate_password_hash(password, method=PASSWORD_METHOD))
    return jsonify({"ok": True, "username": username})


@app.route("/api/users/<username>", methods=["DELETE"])
@require_role("admin")
def api_user_delete(username):
    me = _require_user()
    if username == me["username"]:
        raise ApiError(u"нельзя удалить себя")
    target = db.get_user(username)
    if not target:
        raise ApiError(u"пользователь не найден: %s" % username, 404)
    if target["role"] == "admin" and db.count_role("admin") <= 1:
        raise ApiError(u"нельзя удалить последнего администратора")
    db.delete_user(username)
    return jsonify({"ok": True})


# --------------------------------------------------------------------------- #
# статика
# --------------------------------------------------------------------------- #
@app.route("/")
def index():
    return send_from_directory(app.static_folder, "index.html")


# --------------------------------------------------------------------------- #
# API
# --------------------------------------------------------------------------- #
@app.route("/api/meta")
def api_meta():
    d = data("all")
    meta = d["meta"]
    norms = db.get_norms()
    defs = []
    for m in d.get("metric_defs", []):
        norm = norms.get(m.get("key"), {})
        defs.append({
            "key": m.get("key"),
            "name": m.get("name"),
            "unit": norm.get("unit") or m.get("unit", ""),
            "good_is_higher": norm.get("good_is_higher", bool(m.get("good_is_higher"))),
            "norm": norm.get("value", m.get("norm")),
        })
    return jsonify({
        "company": meta.get("company"),
        "currency": meta.get("currency"),
        "vat_rate": meta.get("vat_rate"),
        "months": meta.get("months"),
        "partial_last_month": meta.get("partial_last_month"),
        "periods": [{"value": p, "label": _period_label(p)} for p in PERIODS],
        "metrics": defs,
        "data_source": meta.get("data_source"),
        "sources": available_sources(),
        "dataset_updated_at": meta.get("dataset_updated_at"),
        "app": APP_NAME,
    })


@app.route("/api/overview")
def api_overview():
    period = parse_period()
    d = data(period)
    dds = d["dds"]
    meta = d["meta"]
    return jsonify({
        "period": period,
        "period_label": meta.get("period_label"),
        "is_partial": meta.get("is_partial"),
        "partial_last_month": meta.get("partial_last_month"),
        "vat_rate": meta.get("vat_rate"),
        "currency": meta.get("currency"),
        "months": dds.get("months"),
        "income": dds.get("income"),
        "expense": dds.get("expense"),
        "income_total": dds.get("income_total"),
        "expense_total": dds.get("expense_total"),
        "cash_result": dds.get("cash_result"),
        "balance": dds.get("balance"),
        "expense_by_article": dds.get("expense_by_article"),
        "expense_by_article_total": dds.get("expense_by_article_total"),
        "expense_full_total": dds.get("expense_full_total"),
        "operations": dds.get("operations"),
    })


@app.route("/api/marketing")
def api_marketing():
    period = parse_period()
    d = data(period)
    mkt = d["marketing"]
    meta = d["meta"]
    return jsonify({
        "period": period,
        "period_label": meta.get("period_label"),
        "is_partial": meta.get("is_partial"),
        "partial_last_month": meta.get("partial_last_month"),
        "vat_rate": meta.get("vat_rate"),
        "months": mkt.get("months"),
        "plan": mkt.get("plan"),
        "fact": mkt.get("fact"),
        "plan_total": mkt.get("plan_total"),
        "fact_total": mkt.get("fact_total"),
        "osvoenie_pct": mkt.get("osvoenie_pct"),
        "deviation_pct": mkt.get("deviation_pct"),
        "channels": mkt.get("channels"),
    })


@app.route("/api/sales")
def api_sales():
    # период к CRM-данным в датасете не применяется (нет разбивки сделок по месяцам)
    d = data("all")
    sales = d["sales"]
    meta = d["meta"]
    return jsonify({
        "period": "all",
        "basis": sales.get("basis", "full"),
        "currency": meta.get("currency"),
        "vat_rate": meta.get("vat_rate"),
        "funnel": sales.get("funnel"),
        "directions": sales.get("directions"),
        "leads": sales.get("leads"),
        "paid_deals": sales.get("paid_deals"),
        "revenue_total": sales.get("revenue_total"),
        "deals_total": sales.get("deals_total"),
        "avg_check": sales.get("avg_check"),
        "conversion_pct": sales.get("conversion_pct"),
    })


def build_metrics(period: str, d: Dict[str, Any]) -> List[Dict[str, Any]]:
    """4 метрики: значение, норматив из SQLite, отклонение, формула, база расчёта."""
    basis = d["basis"]
    per, full = basis["period"], basis["full"]
    norms = db.get_norms()
    # (ключ, числитель, знаменатель, база: period — пересчёт от периода, full — янв–сен)
    calc = [
        ("roas", per["income_net"], per["budget_fact_net"], "period"),
        ("cac", full["budget_fact_net"], basis["paid_deals"], "full"),
        ("gp_per_lead", full["gross_profit"], basis["leads"], "full"),
        ("cash_conv", per["income_net"], per["revenue_accrued_net"], "period"),
    ]
    defs = {m.get("key"): m for m in d.get("metric_defs", [])}

    out = []
    for key, num, den, base in calc:
        meta = defs.get(key, {})
        norm = norms.get(key, {})
        value = num / den if den else 0.0
        norm_value = norm.get("value")
        if norm_value is None:
            norm_value = meta.get("norm", 0.0)
        unit = norm.get("unit") or meta.get("unit", "")
        value = _round_by_unit(value, unit)
        norm_value = _round_by_unit(float(norm_value), unit)
        dev = _deviation(value, norm_value)
        good_is_higher = bool(norm.get("good_is_higher", meta.get("good_is_higher", True)))
        out.append({
            "key": key,
            "name": meta.get("name", key),
            "formula": meta.get("formula", ""),
            "value": value,
            "norm": norm_value,
            "unit": unit,
            "good_is_higher": good_is_higher,
            "basis": base,
            "deviation_pct": dev,
            "is_ok": dev >= 0 if good_is_higher else dev <= 0,
        })
    return out


@app.route("/api/metrics")
def api_metrics():
    period = parse_period()
    d = data(period)
    meta = d["meta"]
    return jsonify({
        "period": period,
        "period_label": meta.get("period_label"),
        "is_partial": meta.get("is_partial"),
        "vat_rate": meta.get("vat_rate"),
        "currency": meta.get("currency"),
        "metrics": build_metrics(period, d),
        "note": (u"CAC и валовая прибыль на лида считаются за весь период датасета "
                 u"(в данных нет разбивки сделок по месяцам), ROAS и Cash Conversion — "
                 u"за выбранный период. Нормативы задаются в Настройках."),
    })


@app.route("/api/settings", methods=["GET", "PUT"])
@require_role("editor")   # настройки и нормативы видят/меняют editor и admin
def api_settings():
    if request.method == "GET":
        return jsonify(_settings_payload())

    payload = request.get_json(silent=True)
    if not isinstance(payload, dict):
        raise ApiError(u"ожидается JSON-объект с ключами settings и/или norms")

    settings_in = payload.get("settings") or {}
    norms_in = payload.get("norms") or {}
    if not isinstance(settings_in, dict) or not isinstance(norms_in, dict):
        raise ApiError(u"settings и norms должны быть объектами")

    unknown = [k for k in settings_in if k not in db.SETTING_KEYS]
    if unknown:
        raise ApiError(u"неизвестные настройки: %s" % ", ".join(sorted(unknown)))

    known_metrics = {m.get("key") for m in data("all").get("metric_defs", [])}
    norms = {}
    for key, value in norms_in.items():
        if key not in known_metrics:
            raise ApiError(u"неизвестный норматив: %s" % key)
        try:
            num = float(str(value).replace(",", ".").replace(" ", ""))
        except (TypeError, ValueError):
            raise ApiError(u"норматив «%s» должен быть числом" % key)
        if num < 0:
            raise ApiError(u"норматив «%s» не может быть отрицательным" % key)
        norms[key] = num

    if settings_in:
        db.save_settings(settings_in)
    if norms:
        db.save_norms(norms)

    body = _settings_payload()
    body["saved"] = True
    body["saved_at"] = db.now()
    body["changed"] = {"settings": sorted(settings_in.keys()), "norms": sorted(norms.keys())}
    return jsonify(body)


def _settings_payload() -> Dict[str, Any]:
    return {
        "settings": db.public_settings(),
        "norms": db.get_norms(),
        "data_source": get_active_source().name,
        "sources": available_sources(),
        "updated_at": max([v.get("updated_at", "") for v in db.get_settings().values()] or [""]),
    }


# --------------------------------------------------------------------------- #
# загрузка ведомости и журнал импортов (Ф5)
# --------------------------------------------------------------------------- #
def _fact_xlsx_path() -> str:
    """Куда сохраняем загруженную ведомость: путь активного источника или data/."""
    src = get_active_source()
    raw = getattr(src, "xlsx_file", None) or os.path.join("data", "fact_marketing.xlsx")
    return raw if os.path.isabs(raw) else os.path.join(BASE_DIR, raw)


def _dataset_shape() -> Any:
    """(месяцы, каналы) из базового датасета — эталон для проверки ведомости.
    Берём из активного источника (он наследует MockSource._dataset), без походов
    в стаб/Google — валидация не должна зависеть от внешних сервисов."""
    src = get_active_source()
    d = src._dataset()
    months = list(d.get("months", []))
    channels = [ch.get("channel", "") for ch in d.get("marketing", {}).get("channels", [])]
    return months, channels


def _parse_fact_workbook(filename: str, blob: bytes) -> Any:
    """Валидация загруженной ведомости по шагам плана. Возвращает (факт, строк).
    Каждая ошибка — ApiError(400) с понятным текстом; файл при этом не заменяется."""
    # 1. расширение
    if not (filename or "").lower().endswith(".xlsx"):
        raise ApiError(u"нужен файл Excel (.xlsx), получено: %s" % (filename or u"без имени"))
    # 2. открываем из потока
    try:
        wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    except Exception:
        raise ApiError(u"не удалось открыть файл: это не похоже на Excel-ведомость")
    ws = wb[SHEET_TITLE] if SHEET_TITLE in wb.sheetnames else wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if len(rows) < 3:  # титул + шапка + хотя бы один канал
        raise ApiError(u"в файле нет шапки и строк с каналами")

    months, expected_channels = _dataset_shape()
    # 3. шапка в строке 2, месяцы совпадают с датасетом
    header = [str(c).strip() if c is not None else "" for c in rows[1]]
    file_months = [m for m in header[1:] if m]
    if not file_months:
        raise ApiError(u"в шапке (строка 2) нет месяцев")
    if file_months != months:
        raise ApiError(u"месяцы в файле (%s) не совпадают с датасетом (%s)"
                       % (u", ".join(file_months), u", ".join(months)))
    # 4. каналы — тот же набор, без дублей
    fact = {}
    for row in rows[2:]:
        if row is None or not row or row[0] is None:
            continue
        channel = str(row[0]).strip()
        if not channel:
            continue
        if channel in fact:
            raise ApiError(u"канал «%s» встречается в файле дважды" % channel)
        values = list(row[1:len(months) + 1])
        if len(values) < len(months):
            values += [None] * (len(months) - len(values))
        parsed = []
        for month, v in zip(months, values):
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                num = float(v)
            else:
                try:
                    num = float(str(v).replace(" ", "").replace(u"\xa0", "").replace(",", "."))
                except (TypeError, ValueError):
                    raise ApiError(u"«%s», %s: значение не число (%r)" % (channel, month, v))
            # 5. значения — числа ≥ 0
            if num < 0:
                raise ApiError(u"«%s», %s: значение не может быть отрицательным" % (channel, month))
            parsed.append(int(round(num)))
        fact[channel] = parsed
    if not fact:
        raise ApiError(u"в файле нет строк с каналами")
    unknown = sorted(set(fact) - set(expected_channels))
    missing = [c for c in expected_channels if c not in fact]
    problems = []
    if unknown:
        problems.append(u"лишние каналы: %s" % u", ".join(unknown))
    if missing:
        problems.append(u"нет каналов: %s" % u", ".join(missing))
    if problems:
        raise ApiError(u"каналы не совпадают с планом (%s)" % u"; ".join(problems))
    return fact, len(fact)


def _current_fact_total(xlsx_path: str) -> float:
    """Сумма факта из текущей ведомости (для контрольной сверки)."""
    reader = ExcelSource(xlsx_file=xlsx_path)
    current, _months = reader._workbook_rows()
    return float(sum(sum(vals) for vals in current.values()))


def _save_fact_atomically(blob: bytes, target: str) -> None:
    """Временный файл рядом с целью + os.replace — замена атомарна."""
    os.makedirs(os.path.dirname(target), exist_ok=True)
    tmp = tempfile.NamedTemporaryFile(
        dir=os.path.dirname(target) or ".", prefix=".upload-", suffix=".xlsx", delete=False)
    try:
        tmp.write(blob)
        tmp.flush()
        tmp.close()
        os.replace(tmp.name, target)
    except Exception:
        tmp.close()
        try:
            os.unlink(tmp.name)
        except OSError:
            pass
        raise


@app.route("/api/upload/fact_marketing", methods=["POST"])
@require_role("editor")
def api_upload_fact_marketing():
    f = request.files.get("file")
    if f is None or not f.filename:
        raise ApiError(u"не передан файл (поле multipart «file»)")
    started = db.now()
    target = _fact_xlsx_path()
    try:
        blob = f.read()
        fact, n_rows = _parse_fact_workbook(f.filename, blob)
        # 6. контрольная сверка: защита от загрузки «не той» ведомости
        try:
            current_total = _current_fact_total(target)
        except SourceError:
            current_total = 0.0   # текущую не прочли — сверку пропускаем
        new_total = float(sum(sum(vals) for vals in fact.values()))
        if current_total > 0 and abs(new_total - current_total) / current_total > FACT_DEVIATION_LIMIT:
            raise ApiError(u"похоже, не та ведомость: сумма в файле (%s) отличается "
                           u"от текущей (%s) больше чем на %.0f%%"
                           % (_fmt_money(new_total), _fmt_money(current_total),
                              FACT_DEVIATION_LIMIT * 100))
        _save_fact_atomically(blob, target)
    except ApiError as exc:
        db.log_import("excel_file", 0, "error", exc.message, started_at=started)
        raise
    db.log_import("excel_file", n_rows, "ok", f.filename, started_at=started)
    src = get_active_source()
    if hasattr(src, "reload"):
        src.reload()          # сбросить кэш ведомости активного источника
    return jsonify({"ok": True, "rows": n_rows, "file": f.filename,
                    "total": int(round(new_total)), "imports": db.recent_imports(5)})


@app.route("/api/imports")
@require_role("editor")
def api_imports():
    return jsonify({"imports": db.recent_imports(20)})


# --------------------------------------------------------------------------- #
# утилиты
# --------------------------------------------------------------------------- #
def _fmt_money(value: float) -> str:
    return u"{:,.0f}".format(value).replace(",", " ")
def _period_label(period: str) -> str:
    return PERIOD_LABELS.get(period, period)


def _round_half_up(value: float, digits: int = 1) -> float:
    """Округление как в JS Math.round — чтобы API и интерфейс не расходились."""
    mult = 10 ** digits
    return math.floor(value * mult + 0.5) / mult


def _round_by_unit(value: float, unit: str) -> float:
    if unit == u"×":
        return round(value, 2)
    if unit == u"%":
        return round(value, 3)
    return round(value)


def _deviation(value: float, norm: float) -> float:
    if not norm:
        return 0.0
    return _round_half_up((value / norm - 1) * 100)


def ensure_admin_user() -> None:
    """Ф5: при первом старте (users пуста) — admin с паролем из .env.
    Пароль в логи не пишем; пользователь меняет его через управление пользователями."""
    if db.count_users() > 0:
        return
    password = os.environ.get("ADMIN_PASSWORD") or ""
    if not password:
        print(u"[users] ADMIN_PASSWORD не задан — администратор не создан")
        return
    db.insert_user("admin", generate_password_hash(password, method=PASSWORD_METHOD), "admin")
    print(u"[users] создан администратор «admin» (пароль — из ADMIN_PASSWORD в .env)")


def init_app() -> str:
    """Создаёт БД, засеивает нормативы, администратора и пишет в журнал импортов."""
    path = db.init_db()
    ensure_admin_user()
    src = get_active_source()
    try:
        rows = src.rows_count()
        db.log_import(src.name, rows, "ok", os.path.basename(getattr(src, "data_file", "mock")))
        print(u"[db] %s · источник: %s · строк: %s" % (path, src.name, rows))
    except SourceError as exc:
        db.log_import(src.name, 0, "error", str(exc))
        print(u"[db] источник недоступен: %s" % exc)
    return path


def main() -> None:
    init_app()
    port = int(os.environ.get("PORT") or 5000)
    debug = (os.environ.get("FLASK_DEBUG") or "1") not in ("0", "false", "False", "")
    print(u"%s -> http://127.0.0.1:%d" % (APP_NAME, port))
    app.run(host="127.0.0.1", port=port, debug=debug)


if __name__ == "__main__":
    main()
