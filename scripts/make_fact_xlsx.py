#!/usr/bin/env python3
"""Excel-ведомость факта маркетинга из тестового датасета (синтетика).

Читает data/fact_marketing.csv (его делает scripts/gen_test_data.py) и собирает
data/fact_marketing.xlsx — файл-источник для источника excel_file (Ф4).

Структура листа «Факт маркетинга»:
  строка 1 — заголовок ведомости, строка 2 — шапка (Канал, месяцы), далее каналы.
Формат намеренно «человеческий»: именно такие ведомости приходят от клиентов.

Запуск:  scripts/make_fact_xlsx.py   (из корня проекта или как есть)
"""
import csv
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CSV_IN = os.path.join(BASE, "data", "fact_marketing.csv")
XLSX_OUT = os.path.join(BASE, "data", "fact_marketing.xlsx")

wb = Workbook()
ws = wb.active
ws.title = u"Факт маркетинга"

with open(CSV_IN, encoding="utf-8") as f:
    rows = list(csv.reader(f))
header, data = rows[0], rows[1:]  # header: Канал, Янв..Сен

bold = Font(bold=True)
band = PatternFill("solid", fgColor="1F2A44")
white_bold = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="D9DDE5")
box = Border(left=thin, right=thin, top=thin, bottom=thin)

ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(header))
c = ws.cell(row=1, column=1, value=u"Ведомость факт маркетинга — 2026 (тестовые данные)")
c.font = Font(bold=True, size=13)
c.alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[1].height = 24

for j, name in enumerate(header, start=1):
    c = ws.cell(row=2, column=j, value=name)
    c.font = white_bold
    c.fill = band
    c.alignment = Alignment(horizontal="center" if j > 1 else "left")
    c.border = box

for i, row in enumerate(data, start=3):
    for j, val in enumerate(row, start=1):
        c = ws.cell(row=i, column=j, value=int(float(val)) if j > 1 else val)
        c.border = box
        if j > 1:
            c.number_format = u"# ##0"

ws.column_dimensions["A"].width = 22
for j in range(2, len(header) + 1):
    ws.column_dimensions[ws.cell(row=2, column=j).column_letter].width = 11
ws.freeze_panes = "B3"

wb.save(XLSX_OUT)
print(u"OK: %s (%d каналов, %d месяцев)" % (os.path.relpath(XLSX_OUT, BASE), len(data), len(header) - 1))
